"""
環境省 一般廃棄物処理実態調査（市区町村別）のExcelデータを SQLite に整形する ETL スクリプト。

data/R1.xlsx 〜 R6.xlsx と data/R{n}cost.xlsx を読み込み、以下を実施する:
  - 「ごみ処理概要」シートから人口・総排出量・処理量・リサイクル率などを抽出
  - 「ごみ搬入量内訳」シートから 生活系(家庭)ごみ搬入量 / 事業系ごみ搬入量 を抽出
  - 市区町村コードで結合し、年度ごとに waste_stats テーブルへ格納
  - 都道府県コードで集計した pref_stats（都道府県別集計）テーブルも生成
  - R{n}cost.xlsx の「廃棄物事業経費（歳出）」からごみ処理事業費（都道府県単位）を
    取得し、pref_stats に処理原価（円/t）として merge
  - 市区町村の緯度経度（localgovjp）を municipality_latlng テーブルへ格納

実行:
  venv/bin/python backend/etl.py
"""
from __future__ import annotations

import csv
import re
import sqlite3
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
GEO_DIR = Path(__file__).resolve().parent / "geo"
DB_PATH = Path(__file__).resolve().parent / "waste.db"

FISCAL_YEARS = ["R1", "R2", "R3", "R4", "R5", "R6"]
FISCAL_YEAR_LABEL = {
    "R1": "令和元年度",
    "R2": "令和2年度",
    "R3": "令和3年度",
    "R4": "令和4年度",
    "R5": "令和5年度",
    "R6": "令和6年度",
}
# Rows to exclude from prefecture-level sums and municipality-level rankings/maps:
#   48000 = 全国合計 (national total, not a municipality)
#   13101-13123 = the 23 special wards of Tokyo. Population is reported per-ward, but
#     all actual waste figures (収集量/処理量/リサイクル率 etc.) are reported ONLY as a
#     single combined row "東京都23区" (13100). Counting both would double the population
#     and wrongly zero out the waste metrics for the wards, so the 23 individual ward
#     rows are excluded here and 13100 is kept as the municipality-level unit for them.
NATIONAL_TOTAL_CODE = "48000"
TOKYO_WARD_CODES = {f"131{n:02d}" for n in range(1, 24)}  # 13101-13123
TOKYO_23_WARDS_COMBINED_CODE = "13100"
AGGREGATE_CODES = {NATIONAL_TOTAL_CODE} | TOKYO_WARD_CODES


def _norm(v):
    if v is None:
        return None
    s = str(v).replace("\n", "").replace("　", "").strip()
    return s or None


def read_header_matrix(ws, header_rows=(2, 3, 4)):
    """Return list of (r2, r3, r4) tuples per 0-indexed column, raw (no forward fill)."""
    rows = list(ws.iter_rows(min_row=header_rows[0], max_row=header_rows[-1], values_only=True))
    ncol = ws.max_column
    matrix = []
    for i in range(ncol):
        vals = [_norm(rows[r][i]) if i < len(rows[r]) else None for r in range(len(rows))]
        matrix.append(tuple(vals))
    return matrix


def forward_fill(seq):
    out = []
    last = None
    for v in seq:
        if v is not None:
            last = v
        out.append(last)
    return out


def find_column(
    matrix,
    *,
    group_prefix=None,
    group_exclude=None,
    sub_prefix=None,
    leaf_exact=None,
    leaf_prefix=None,
    require_no_r4=False,
):
    """Locate a single column index in the (r2, r3, r4) header matrix.

    group_prefix: substring the forward-filled r2 (group label) must start with
    sub_prefix: substring the forward-filled r3 (subgroup label) must start with
        (needed to tell one subgroup's "合計" from another's within the same group)
    leaf_exact / leaf_prefix: match against leaf label = first of r4, r3, r2 that is not None
    require_no_r4: reject columns whose leaf comes from a nested r4 sub-breakdown
        (needed to disambiguate a subgroup's own "合計" from the parent group's overall "合計")
    """
    r2_filled = forward_fill([m[0] for m in matrix])
    r3_filled = forward_fill([m[1] for m in matrix])
    candidates = []
    for i, m in enumerate(matrix):
        r2, r3, r4 = m[0], m[1], m[2] if len(m) > 2 else None
        leaf = r4 or r3 or r2
        group = r2_filled[i]
        if group_prefix is not None:
            if group is None or not group.startswith(group_prefix):
                continue
        if sub_prefix is not None:
            sub = r3_filled[i]
            if sub is None or not sub.startswith(sub_prefix):
                continue
        if group_exclude is not None:
            if group and group_exclude in group:
                continue
        if require_no_r4 and r4 is not None:
            continue
        if leaf_exact is not None and leaf != leaf_exact:
            continue
        if leaf_prefix is not None and (leaf is None or not leaf.startswith(leaf_prefix)):
            continue
        candidates.append(i)
    if not candidates:
        return None
    return candidates[0]


def to_number(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "-", "―", "*", "…"):
        return None
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def parse_overview_sheet(ws):
    matrix = read_header_matrix(ws)
    col = {
        "pref_name": find_column(matrix, leaf_exact="都道府県名"),
        "city_code": find_column(matrix, leaf_exact="地方公共団体コード"),
        "city_name": find_column(matrix, leaf_exact="市区町村名"),
        "population": find_column(matrix, leaf_exact="総人口"),
        "total_waste_t": find_column(matrix, group_prefix="ごみ総排出量", leaf_exact="合計"),
        "per_capita_total_g": find_column(matrix, group_prefix="１人１日当たりの排出量", leaf_prefix="合計"),
        "per_capita_household_g": find_column(matrix, group_prefix="１人１日当たりの排出量", leaf_prefix="生活系ごみ"),
        "per_capita_business_g": find_column(matrix, group_prefix="１人１日当たりの排出量", leaf_prefix="事業系ごみ"),
        "treated_amount_t": find_column(matrix, group_prefix="ごみ処理量", leaf_exact="合計", require_no_r4=True),
        "reduction_rate_pct": find_column(matrix, group_prefix="減量処理率"),
        "recycling_rate_r_pct": find_column(
            matrix, group_prefix="リサイクル率 Ｒ", group_exclude="Ｒ’"
        ),
        "recycling_rate_r2_pct": find_column(matrix, group_prefix="リサイクル率 Ｒ’"),
        "final_disposal_t": find_column(matrix, group_prefix="最終処分量", leaf_exact="合計"),
    }
    missing = [k for k, v in col.items() if v is None]
    if missing:
        raise RuntimeError(f"overview: columns not found: {missing}")

    records = {}
    data_start_row = 7
    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        code = _norm(row[col["city_code"]])
        if not code:
            continue
        rec = {
            "pref_name": _norm(row[col["pref_name"]]),
            "city_code": code,
            "city_name": _norm(row[col["city_name"]]),
            "population": to_number(row[col["population"]]),
            "total_waste_t": to_number(row[col["total_waste_t"]]),
            "per_capita_total_g": to_number(row[col["per_capita_total_g"]]),
            "per_capita_household_g": to_number(row[col["per_capita_household_g"]]),
            "per_capita_business_g": to_number(row[col["per_capita_business_g"]]),
            "treated_amount_t": to_number(row[col["treated_amount_t"]]),
            "reduction_rate_pct": to_number(row[col["reduction_rate_pct"]]),
            "recycling_rate_r_pct": to_number(row[col["recycling_rate_r_pct"]]),
            "recycling_rate_r2_pct": to_number(row[col["recycling_rate_r2_pct"]]),
            "final_disposal_t": to_number(row[col["final_disposal_t"]]),
        }
        records[code] = rec
    return records


def parse_inflow_sheet(ws):
    matrix = read_header_matrix(ws)
    col = {
        "city_code": find_column(matrix, leaf_exact="地方公共団体コード"),
        "household_waste_t": find_column(matrix, leaf_exact="生活系ごみ搬入量"),
        "business_waste_t": find_column(matrix, leaf_exact="事業系ごみ搬入量"),
        "self_treatment_t": find_column(matrix, leaf_exact="自家処理量"),
    }
    missing = [k for k, v in col.items() if v is None]
    if missing:
        raise RuntimeError(f"inflow: columns not found: {missing}")

    records = {}
    for row in ws.iter_rows(min_row=7, values_only=True):
        code = _norm(row[col["city_code"]])
        if not code:
            continue
        records[code] = {
            "household_waste_t": to_number(row[col["household_waste_t"]]),
            "business_waste_t": to_number(row[col["business_waste_t"]]),
            "self_treatment_t": to_number(row[col["self_treatment_t"]]),
        }
    return records


def load_workbook_year(year: str):
    path = DATA_DIR / f"{year}.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    overview = parse_overview_sheet(wb["ごみ処理概要"])
    inflow = parse_inflow_sheet(wb["ごみ搬入量内訳"])
    wb.close()

    # In R1-R4, the "東京都23区" combined row (13100) that carries the actual waste
    # figures has population=0 (a gap in the source data); back-fill it from the sum
    # of the 23 individual ward rows, which always have real population but no waste
    # data of their own (see AGGREGATE_CODES comment above).
    combined = overview.get(TOKYO_23_WARDS_COMBINED_CODE)
    if combined and not combined.get("population"):
        ward_pop_sum = sum(
            overview[c]["population"]
            for c in TOKYO_WARD_CODES
            if c in overview and overview[c].get("population")
        )
        if ward_pop_sum:
            combined["population"] = ward_pop_sum

    merged = []
    for code, rec in overview.items():
        extra = inflow.get(code, {})
        row = {**rec, **extra, "fiscal_year": year, "is_aggregate": 1 if code in AGGREGATE_CODES else 0}
        row["pref_code"] = code[:2] if code and code != "48000" else None
        merged.append(row)
    return merged


# 「廃棄物事業経費」ファイル（R{n}cost.xlsx）は都道府県単位の集計のみで、
# 市区町村別のコストは含まれていない。
COST_FISCAL_YEARS = ["R1", "R2", "R3", "R4", "R5", "R6"]


def parse_cost_sheet(ws):
    """Extract prefecture-level ごみ(waste, excl. し尿) expenditure total (千円) from
    the 廃棄物事業経費（歳出） sheet of an R{n}cost.xlsx workbook."""
    matrix = read_header_matrix(ws)
    col = {
        "pref_name": find_column(matrix, leaf_exact="都道府県名"),
        "city_code": find_column(matrix, leaf_exact="地方公共団体コード"),
        # ごみ分の歳出総額（建設改良費＋処理及び維持管理費＋その他）
        "waste_expenditure_thousand_yen": find_column(
            matrix, group_prefix="ごみ（建設改良費", leaf_exact="合計", require_no_r4=True
        ),
        # その内訳。施設整備の年だけ原価が跳ね上がるため、両者を分けて保持する。
        "construction_expenditure_thousand_yen": find_column(
            matrix, group_prefix="ごみ（建設改良費", sub_prefix="建設改良費", leaf_exact="合計"
        ),
        "operating_expenditure_thousand_yen": find_column(
            matrix, group_prefix="ごみ（建設改良費", sub_prefix="処理及び維持管理費", leaf_exact="合計"
        ),
    }
    missing = [k for k, v in col.items() if v is None]
    if missing:
        raise RuntimeError(f"cost: columns not found: {missing}")

    records = {}
    for row in ws.iter_rows(min_row=7, values_only=True):
        code = _norm(row[col["city_code"]])
        if not code or not code.endswith("000"):
            continue
        pref_code = code[:2]
        records[pref_code] = {
            "pref_name": _norm(row[col["pref_name"]]),
            "waste_expenditure_thousand_yen": to_number(row[col["waste_expenditure_thousand_yen"]]),
            "construction_expenditure_thousand_yen": to_number(
                row[col["construction_expenditure_thousand_yen"]]
            ),
            "operating_expenditure_thousand_yen": to_number(
                row[col["operating_expenditure_thousand_yen"]]
            ),
        }

    # 内訳（建設改良費・処理及び維持管理費）は必ず総額の一部分になる。列の解決を
    # 誤ると内訳が総額を超えるので、レイアウトが変わったときにここで気付ける。
    for pref_code, rec in records.items():
        total = rec["waste_expenditure_thousand_yen"]
        parts = [rec["construction_expenditure_thousand_yen"], rec["operating_expenditure_thousand_yen"]]
        if total is None or any(p is None for p in parts):
            continue
        if sum(parts) > total + 1:
            raise RuntimeError(
                f"cost: 内訳が総額を超えています（列の解決ミスの可能性） pref={pref_code} "
                f"建設={parts[0]} 運営={parts[1]} 合計={total}"
            )
    return records


def load_cost_year(year: str):
    path = DATA_DIR / f"{year}cost.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = parse_cost_sheet(wb["廃棄物事業経費（歳出）"])
    wb.close()
    return records


# 歳出の3区分のうち、原価に換算して保持するもの。建設改良費は施設を建て替えた
# 年度にだけ集中して計上されるため、合計原価だけを見ると「その年に工事をしていた
# か」で順位が入れ替わってしまう。運営費（処理及び維持管理費）を分けて持つことで、
# 日々の運営にかかるコストだけを年度またぎで比較できるようにする。
COST_BREAKDOWN = [
    ("waste_expenditure_thousand_yen", "cost_per_ton_yen"),
    ("construction_expenditure_thousand_yen", "construction_cost_per_ton_yen"),
    ("operating_expenditure_thousand_yen", "operating_cost_per_ton_yen"),
]


def build_cost_metrics(conn):
    """Merge per-prefecture waste expenditure into pref_stats and derive 処理原価
    (cost per tonne treated), both in total and split into 建設改良費 / 運営費.
    Runs after build_pref_stats() so treated_amount_t is already populated for
    every fiscal_year/pref_code."""
    cur = conn.cursor()

    exp_cols = [e for e, _ in COST_BREAKDOWN]
    rate_cols = [r for _, r in COST_BREAKDOWN]
    set_clause = ", ".join(f"{c} = ?" for c in exp_cols + rate_cols)

    all_period_expenditure: dict[str, dict[str, float]] = {}
    all_period_treated: dict[str, float] = {}

    for year in COST_FISCAL_YEARS:
        cost = load_cost_year(year)
        for pref_code, rec in cost.items():
            if rec["waste_expenditure_thousand_yen"] is None:
                continue
            row = cur.execute(
                "SELECT treated_amount_t FROM pref_stats WHERE fiscal_year = ? AND pref_code = ?",
                (year, pref_code),
            ).fetchone()
            treated = row[0] if row else None

            exps = [rec[c] for c in exp_cols]
            rates = [(e * 1000 / treated) if (e is not None and treated) else None for e in exps]
            cur.execute(
                f"UPDATE pref_stats SET {set_clause} WHERE fiscal_year = ? AND pref_code = ?",
                (*exps, *rates, year, pref_code),
            )

            sums = all_period_expenditure.setdefault(pref_code, {c: 0.0 for c in exp_cols})
            for c, e in zip(exp_cols, exps):
                sums[c] += e or 0.0
            if treated:
                all_period_treated[pref_code] = all_period_treated.get(pref_code, 0.0) + treated

    for pref_code, sums in all_period_expenditure.items():
        treated_sum = all_period_treated.get(pref_code)
        exps = [sums[c] for c in exp_cols]
        rates = [(e * 1000 / treated_sum) if treated_sum else None for e in exps]
        cur.execute(
            f"UPDATE pref_stats SET {set_clause} WHERE fiscal_year = 'ALL' AND pref_code = ?",
            (*exps, *rates, pref_code),
        )
    conn.commit()


SCHEMA = """
DROP TABLE IF EXISTS waste_stats;
CREATE TABLE waste_stats (
    fiscal_year TEXT NOT NULL,
    pref_code TEXT,
    pref_name TEXT,
    city_code TEXT NOT NULL,
    city_name TEXT,
    is_aggregate INTEGER NOT NULL DEFAULT 0,
    population REAL,
    total_waste_t REAL,
    household_waste_t REAL,
    business_waste_t REAL,
    self_treatment_t REAL,
    per_capita_total_g REAL,
    per_capita_household_g REAL,
    per_capita_business_g REAL,
    treated_amount_t REAL,
    reduction_rate_pct REAL,
    recycling_rate_r_pct REAL,
    recycling_rate_r2_pct REAL,
    final_disposal_t REAL,
    PRIMARY KEY (fiscal_year, city_code)
);
CREATE INDEX idx_waste_stats_pref ON waste_stats(fiscal_year, pref_code);
CREATE INDEX idx_waste_stats_year ON waste_stats(fiscal_year);

DROP TABLE IF EXISTS pref_stats;
CREATE TABLE pref_stats (
    fiscal_year TEXT NOT NULL,
    pref_code TEXT NOT NULL,
    pref_name TEXT,
    population REAL,
    total_waste_t REAL,
    household_waste_t REAL,
    business_waste_t REAL,
    treated_amount_t REAL,
    final_disposal_t REAL,
    recycled_t REAL,
    per_capita_total_g REAL,
    per_capita_household_g REAL,
    per_capita_business_g REAL,
    recycling_rate_r_pct REAL,
    final_disposal_rate_pct REAL,
    waste_expenditure_thousand_yen REAL,
    cost_per_ton_yen REAL,
    construction_expenditure_thousand_yen REAL,
    construction_cost_per_ton_yen REAL,
    operating_expenditure_thousand_yen REAL,
    operating_cost_per_ton_yen REAL,
    PRIMARY KEY (fiscal_year, pref_code)
);

DROP TABLE IF EXISTS municipality_latlng;
CREATE TABLE municipality_latlng (
    city_code TEXT PRIMARY KEY,
    pref_name TEXT,
    city_name TEXT,
    lat REAL,
    lng REAL
);

DROP TABLE IF EXISTS fiscal_years;
CREATE TABLE fiscal_years (
    fiscal_year TEXT PRIMARY KEY,
    label TEXT,
    sort_order INTEGER
);
"""


def build_pref_stats(conn):
    """Aggregate municipality-level rows into prefecture totals (recompute rates, don't average them)."""
    cur = conn.cursor()
    cur.execute(
        """
        SELECT fiscal_year, pref_code, pref_name,
               SUM(population), SUM(total_waste_t), SUM(household_waste_t), SUM(business_waste_t),
               SUM(treated_amount_t), SUM(final_disposal_t),
               SUM(treated_amount_t * recycling_rate_r_pct / 100.0)
        FROM waste_stats
        WHERE is_aggregate = 0 AND pref_code IS NOT NULL
        GROUP BY fiscal_year, pref_code
        """
    )
    rows = cur.fetchall()
    out = []
    for (fy, pref_code, pref_name, pop, total_w, hh_w, biz_w, treated, final_disp, recycled_est) in rows:
        # 'ALL' aggregates flow totals over every fiscal year, so the per-capita
        # denominator needs that many years' worth of days, not a single year's 365.
        days = 365 * len(FISCAL_YEARS) if fy == "ALL" else 365
        per_capita_total_g = (total_w * 1_000_000 / pop / days) if (total_w and pop) else None
        per_capita_hh_g = (hh_w * 1_000_000 / pop / days) if (hh_w and pop) else None
        per_capita_biz_g = (biz_w * 1_000_000 / pop / days) if (biz_w and pop) else None
        recycling_rate = (recycled_est / treated * 100.0) if (recycled_est is not None and treated) else None
        final_rate = (final_disp / treated * 100.0) if (final_disp is not None and treated) else None
        out.append(
            (
                fy, pref_code, pref_name, pop, total_w, hh_w, biz_w, treated, final_disp, recycled_est,
                per_capita_total_g, per_capita_hh_g, per_capita_biz_g, recycling_rate, final_rate,
            )
        )
    cur.executemany(
        """
        INSERT INTO pref_stats (
            fiscal_year, pref_code, pref_name, population, total_waste_t, household_waste_t, business_waste_t,
            treated_amount_t, final_disposal_t, recycled_t,
            per_capita_total_g, per_capita_household_g, per_capita_business_g,
            recycling_rate_r_pct, final_disposal_rate_pct
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        out,
    )
    conn.commit()


def load_municipality_latlng(conn):
    """Load municipality centroid lat/lng from localgovjp CSV, matched to our 5-digit city codes."""
    csv_path = GEO_DIR / "municipality_latlng.csv"
    cur = conn.cursor()
    cur.execute("SELECT DISTINCT city_code, pref_name, city_name FROM waste_stats WHERE is_aggregate = 0")
    known_codes = {row[0]: (row[1], row[2]) for row in cur.fetchall()}

    rows = []
    with open(csv_path, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            lgcode = (r.get("lgcode") or "").strip()
            code5 = lgcode[:5]
            if code5 in known_codes and code5 not in {x[0] for x in rows}:
                lat = to_number(r.get("lat"))
                lng = to_number(r.get("lng"))
                if lat is None or lng is None:
                    continue
                pref_name, city_name = known_codes[code5]
                rows.append((code5, pref_name, city_name, lat, lng))

    if TOKYO_23_WARDS_COMBINED_CODE in known_codes and TOKYO_23_WARDS_COMBINED_CODE not in {x[0] for x in rows}:
        # localgovjp has no single entry for "東京都23区"; use the unweighted centroid
        # of the 23 individual ward coordinates (which the CSV does have).
        ward_coords = []
        with open(csv_path, encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for r in reader:
                if (r.get("lgcode") or "").strip()[:5] in TOKYO_WARD_CODES:
                    lat, lng = to_number(r.get("lat")), to_number(r.get("lng"))
                    if lat is not None and lng is not None:
                        ward_coords.append((lat, lng))
        if ward_coords:
            avg_lat = sum(c[0] for c in ward_coords) / len(ward_coords)
            avg_lng = sum(c[1] for c in ward_coords) / len(ward_coords)
            pref_name, city_name = known_codes[TOKYO_23_WARDS_COMBINED_CODE]
            rows.append((TOKYO_23_WARDS_COMBINED_CODE, pref_name, city_name, avg_lat, avg_lng))

    cur.executemany(
        "INSERT OR IGNORE INTO municipality_latlng (city_code, pref_name, city_name, lat, lng) VALUES (?, ?, ?, ?, ?)",
        rows,
    )
    conn.commit()
    return len(rows), len(known_codes)


WASTE_STATS_INSERT_SQL = """
    INSERT OR REPLACE INTO waste_stats (
        fiscal_year, pref_code, pref_name, city_code, city_name, is_aggregate,
        population, total_waste_t, household_waste_t, business_waste_t, self_treatment_t,
        per_capita_total_g, per_capita_household_g, per_capita_business_g,
        treated_amount_t, reduction_rate_pct, recycling_rate_r_pct, recycling_rate_r2_pct, final_disposal_t
    ) VALUES (
        :fiscal_year, :pref_code, :pref_name, :city_code, :city_name, :is_aggregate,
        :population, :total_waste_t, :household_waste_t, :business_waste_t, :self_treatment_t,
        :per_capita_total_g, :per_capita_household_g, :per_capita_business_g,
        :treated_amount_t, :reduction_rate_pct, :recycling_rate_r_pct, :recycling_rate_r2_pct, :final_disposal_t
    )
"""


def build_all_period_municipalities(conn):
    """Create a synthetic fiscal_year='ALL' row per municipality summarizing R1-R6.

    Flow quantities (tonnage) are summed across available years. Population is the
    latest year's snapshot (a stock, not a flow, so it is not summed). Per-capita
    rates are recomputed as population-day-weighted averages, and percentage rates
    (recycling / reduction) as treated-amount-weighted averages, rather than naively
    averaging percentages or summing them.
    """
    cur = conn.cursor()
    cur.execute(
        """
        SELECT city_code, fiscal_year, pref_code, pref_name, city_name, is_aggregate,
               population, total_waste_t, household_waste_t, business_waste_t, self_treatment_t,
               treated_amount_t, final_disposal_t,
               reduction_rate_pct, recycling_rate_r_pct, recycling_rate_r2_pct
        FROM waste_stats
        WHERE fiscal_year != 'ALL'
        ORDER BY city_code, fiscal_year
        """
    )
    by_city = {}
    for row in cur.fetchall():
        by_city.setdefault(row[0], []).append(row)

    def sum_col(rows, idx):
        vals = [r[idx] for r in rows if r[idx] is not None]
        return sum(vals) if vals else None

    def per_capita(rows, flow_idx):
        num = den = 0.0
        has = False
        for r in rows:
            flow, pop = r[flow_idx], r[6]
            if flow is not None and pop:
                num += flow * 1_000_000
                den += pop * 365
                has = True
        return (num / den) if has and den else None

    def weighted_rate(rows, rate_idx, weight_idx=11):
        num = den = 0.0
        has = False
        for r in rows:
            rate, w = r[rate_idx], r[weight_idx]
            if rate is not None and w:
                num += rate * w
                den += w
                has = True
        return (num / den) if has and den else None

    out = []
    for city_code, rows in by_city.items():
        latest = sorted(rows, key=lambda r: r[1])[-1]
        out.append(
            {
                "fiscal_year": "ALL",
                "pref_code": latest[2],
                "pref_name": latest[3],
                "city_code": city_code,
                "city_name": latest[4],
                "is_aggregate": latest[5],
                "population": latest[6],
                "total_waste_t": sum_col(rows, 7),
                "household_waste_t": sum_col(rows, 8),
                "business_waste_t": sum_col(rows, 9),
                "self_treatment_t": sum_col(rows, 10),
                "per_capita_total_g": per_capita(rows, 7),
                "per_capita_household_g": per_capita(rows, 8),
                "per_capita_business_g": per_capita(rows, 9),
                "treated_amount_t": sum_col(rows, 11),
                "reduction_rate_pct": weighted_rate(rows, 13),
                "recycling_rate_r_pct": weighted_rate(rows, 14),
                "recycling_rate_r2_pct": weighted_rate(rows, 15),
                "final_disposal_t": sum_col(rows, 12),
            }
        )

    cur.executemany(WASTE_STATS_INSERT_SQL, out)
    conn.commit()
    return len(out)


def main():
    if DB_PATH.exists():
        DB_PATH.unlink()
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(SCHEMA)

    cur = conn.cursor()
    for i, year in enumerate(FISCAL_YEARS):
        print(f"[ETL] loading {year} ...")
        rows = load_workbook_year(year)
        cur.executemany(WASTE_STATS_INSERT_SQL, rows)
        cur.execute(
            "INSERT OR REPLACE INTO fiscal_years (fiscal_year, label, sort_order) VALUES (?, ?, ?)",
            (year, FISCAL_YEAR_LABEL[year], i),
        )
        conn.commit()
        print(f"[ETL]   {len(rows)} municipality rows")

    print("[ETL] building ALL-period (R1-R6 combined) municipality rows ...")
    n_all = build_all_period_municipalities(conn)
    cur.execute(
        "INSERT OR REPLACE INTO fiscal_years (fiscal_year, label, sort_order) VALUES (?, ?, ?)",
        ("ALL", "全期間合計（令和元〜6年度）", len(FISCAL_YEARS)),
    )
    conn.commit()
    print(f"[ETL]   {n_all} municipality rows")

    print("[ETL] building pref_stats aggregates ...")
    build_pref_stats(conn)

    print("[ETL] merging waste expenditure cost data (R1-R6) ...")
    build_cost_metrics(conn)

    print("[ETL] loading municipality lat/lng ...")
    matched, total = load_municipality_latlng(conn)
    print(f"[ETL]   matched {matched}/{total} municipalities to coordinates")

    conn.close()
    print(f"[ETL] done -> {DB_PATH}")


if __name__ == "__main__":
    main()
